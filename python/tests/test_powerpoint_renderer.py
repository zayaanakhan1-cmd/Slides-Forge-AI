"""PowerPoint renderer verification.

Proves that the canonical presentation model can be rendered to a real .pptx
file. The test does not mock python-pptx: it renders a model, opens the result
with python-pptx and asserts the file is a valid presentation containing the
expected slide content.
"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation as PptxPresentation

from python.destinations.powerpoint import PowerPointRenderer, PowerPointRenderError


def _model() -> dict:
    return {
        "id": "pres_test",
        "title": "How mitosis works",
        "theme": {
            "colors": {
                "background": "#0B0D12",
                "text": "#F4F6FB",
                "accent": "#6D8BFF",
            },
            "typography": {"scale": {"title": 40, "body": 18}},
        },
        "slides": [
            {
                "id": "slide_1",
                "title": "How mitosis works",
                "metadata": {"dimensions": {"width": 960, "height": 540}},
                "elements": [
                    {
                        "id": "el_1",
                        "type": "text",
                        "position": {"x": 0.1, "y": 0.3},
                        "size": {"width": 0.8, "height": 0.3},
                        "runs": [
                            {"text": "Mitosis produces two identical ", "style": "body"},
                            {"text": "daughter cells", "bold": True},
                        ],
                    }
                ],
            }
        ],
    }


def test_render_produces_a_valid_pptx() -> None:
    renderer = PowerPointRenderer()
    data = renderer.render(_model())

    assert data[:2] == b"PK", "A .pptx is a zip archive and must start with PK"
    assert len(data) > 0

    reopened = PptxPresentation(io.BytesIO(data))
    assert len(reopened.slides) == 1

    texts = []
    for shape in reopened.slides[0].shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)

    joined = "\n".join(texts)
    assert "How mitosis works" in joined
    assert "daughter cells" in joined


def test_render_preserves_slide_count() -> None:
    model = _model()
    model["slides"].append(
        {
            "id": "slide_2",
            "title": "Second slide",
            "elements": [],
        }
    )

    data = PowerPointRenderer().render(model)
    reopened = PptxPresentation(io.BytesIO(data))
    assert len(reopened.slides) == 2


def test_render_rejects_missing_slides() -> None:
    with pytest.raises(PowerPointRenderError):
        PowerPointRenderer().render({"id": "pres_bad"})
